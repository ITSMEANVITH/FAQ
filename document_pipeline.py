
import pymupdf
from docx import Document
from io import BytesIO
from pathlib import Path


# ---------------- DOCUMENT EXTRACTION ----------------

def extract_document(uploaded_file):
    filename = uploaded_file.name
    extension = Path(filename).suffix.lower()
    file_bytes = uploaded_file.getvalue()

    pages = []

    # Extract text from PDF
    if extension == ".pdf":
        pdf = pymupdf.open(
            stream=file_bytes,
            filetype="pdf"
        )

        try:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()

                if text:
                    pages.append({
                        "source": filename,
                        "page": page_number,
                        "text": text
                    })
        finally:
            pdf.close()

    # Extract text from Word document
    elif extension == ".docx":
        doc = Document(BytesIO(file_bytes))
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())

        for table in doc.tables:
            for row in table.rows:
                text_parts.append(
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                )

        text = "\n".join(text_parts).strip()

        if text:
            pages.append({
                "source": filename,
                "page": None,
                "text": text
            })

    # Extract text from TXT file
    elif extension == ".txt":
        text = file_bytes.decode(
            "utf-8-sig",
            errors="replace"
        ).strip()

        if text:
            pages.append({
                "source": filename,
                "page": None,
                "text": text
            })

    else:
        raise ValueError(
            "Upload a PDF, DOCX or TXT file."
        )

    if not pages:
        raise ValueError(
            "No readable text found in this document. "
            "The PDF may be scanned or image-based."
        )

    return pages


# ---------------- TEXT CHUNKING ----------------

def split_into_chunks(pages, chunk_size=500, overlap=80):
    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero.")

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError(
            "overlap must be non-negative and smaller than chunk_size."
        )

    chunks = []
    chunk_id = 0

    for page in pages:
        text = page["text"].strip()

        if not text:
            continue

        start = 0

        while start < len(text):
            end = min(start + chunk_size, len(text))

            # Prefer ending at a sentence boundary.
            if end < len(text):
                sentence_end = max(
                    text.rfind(". ", start, end),
                    text.rfind("? ", start, end),
                    text.rfind("! ", start, end)
                )

                # Avoid creating a very short chunk.
                if sentence_end >= start + chunk_size // 2:
                    end = sentence_end + 1

                else:
                    # Otherwise, avoid cutting a word.
                    space = text.rfind(" ", start, end)

                    if space > start:
                        end = space

            chunk_text = text[start:end].strip()

            if chunk_text:
                chunks.append({
                    "chunk_id": chunk_id,
                    "source": page["source"],
                    "page": page["page"],
                    "text": chunk_text
                })

                chunk_id += 1

            if end >= len(text):
                break

            # Keep overlapping text for context.
            start = max(start + 1, end - overlap)

    return chunks
